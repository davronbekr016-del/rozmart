"""Импорт каталога MVP в базу: штучные позиции ROZMETOV.

Источники:
  • «Сопоставление товаров REGOS — витрина.xlsx», лист «Каталог MVP» — состав, названия
  • «Выгрузка номенклатуры REGOS (исходная).xlsx» — цены (средняя цена продажи)
  • папка прототипа — файлы фотографий, именование p-{код}.jpg

Запуск из корня проекта:  .venv\\Scripts\\python.exe -m scripts.import_catalog

Повторный запуск безопасен. Товар опознаётся по коду REGOS, а не по названию,
поэтому правка витринного названия или смена категории обновляют существующую
карточку и не плодят дубли.
"""
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.db import SessionLocal, init_db
from app.models import Category, Product, Variant

DOCS = Path(r"C:\Users\User\Downloads\Проект ROZMART\1. Документы")
CATALOG_XLSX = DOCS / "Сопоставление товаров REGOS — витрина.xlsx"
REGOS_XLSX = DOCS / "Выгрузка номенклатуры REGOS (исходная).xlsx"
PHOTO_SRC = Path(r"C:\Users\User\Downloads\Проект ROZMART\3. Прототипы\ROZMART_prototype\img")
PHOTO_DST = Path(__file__).resolve().parent.parent / "static" / "img"

CATEGORY_ORDER = ["Копчёная колбаса", "Варёная колбаса", "Сосиски", "Деликатесы"]
CATEGORY_LABEL = {
    "Копчёная колбаса": "Копчёные колбасы",
    "Варёная колбаса": "Варёные колбасы",
    "Сосиски": "Сосиски",
    "Деликатесы": "Деликатесы",
}

# ожидаемые заголовки: выгрузки — это отчёты, вставка колонки сдвинет чтение
# и цены подставятся молча из соседней колонки
CATALOG_COLUMNS = {2: "Код REGOS", 3: "Наименование в REGOS", 4: "Группа REGOS",
                   6: "Название для клиента", 7: "Фасовка"}
REGOS_COLUMNS = {2: "Код номенкл.", 17: "Сред. цена за ед."}


class ImportError_(Exception):
    """Ошибка входных данных, при которой импорт продолжать нельзя."""


def check_columns(ws, expected: dict[int, str], header_row: int, source: str) -> None:
    for col, must_contain in expected.items():
        actual = str(ws.cell(row=header_row, column=col).value or "").strip()
        if must_contain.lower() not in actual.lower():
            raise ImportError_(
                f"{source}: в колонке {col} ожидался заголовок «{must_contain}», "
                f"а найдено «{actual}». Похоже, структура файла изменилась."
            )


def weight_to_grams(weight: str) -> int:
    """«455 г» -> 455, «1 кг» -> 1000. Нужно для сортировки фасовок в карточке."""
    m = re.search(r"(\d+[.,]?\d*)\s*(кг|г)", weight or "")
    if not m:
        return 0
    value = float(m.group(1).replace(",", "."))
    return round(value * 1000) if m.group(2) == "кг" else round(value)


def load_prices() -> dict[str, int]:
    """Средняя цена продажи из отчёта REGOS."""
    ws = load_workbook(REGOS_XLSX, data_only=True)["Анализ продаж"]
    check_columns(ws, REGOS_COLUMNS, header_row=2, source="Выгрузка REGOS")
    prices: dict[str, int] = {}
    for row in range(3, ws.max_row + 1):
        code = ws.cell(row=row, column=2).value
        raw = ws.cell(row=row, column=17).value
        if not code or raw is None:
            continue
        try:
            value = round(float(raw))
        except (TypeError, ValueError):
            print(f"  строка {row}: цена «{raw}» не число, позиция останется без цены")
            continue
        # в отчёте у непродававшихся позиций стоит 0 — это «цена неизвестна»,
        # а не «бесплатно»: такую фасовку публиковать нельзя (BR-37)
        if value > 0:
            prices[str(code).strip()] = value
    return prices


def load_catalog() -> list[dict]:
    """Позиции листа «Каталог MVP»."""
    ws = load_workbook(CATALOG_XLSX, data_only=True)["Каталог MVP"]
    check_columns(ws, CATALOG_COLUMNS, header_row=1, source="Каталог MVP")
    items, seen = [], set()
    for row in range(2, ws.max_row + 1):
        code = ws.cell(row=row, column=2).value
        if not code:
            continue
        code = str(code).strip()
        name = str(ws.cell(row=row, column=6).value or "").strip()
        if not name:
            print(f"  строка {row}: у позиции {code} пустое название — пропущена")
            continue
        if code in seen:
            print(f"  строка {row}: код {code} встречается повторно — пропущена")
            continue
        seen.add(code)
        regos_name = str(ws.cell(row=row, column=3).value or "").upper()
        # вакуумная упаковка выглядит иначе — это отдельная витринная карточка
        if "(VAK)" in regos_name or "V/U" in regos_name:
            name = f"{name} в вакууме"
        items.append({
            "code": code,
            "group": str(ws.cell(row=row, column=4).value or "").strip(),
            "name": name,
            "weight": str(ws.cell(row=row, column=7).value or "").strip(),
        })
    return items


def copy_photo(code: str, product_id: int) -> str | None:
    """Копирует фото в static/img под номером товара. Возвращает имя файла.

    Имя файла — номер товара, а не код REGOS: адрес картинки виден клиенту,
    а служебные коды учётной системы наружу отдавать нельзя (BR-36).
    """
    src = PHOTO_SRC / f"p-{code}.jpg"
    if not src.exists():
        return None
    PHOTO_DST.mkdir(parents=True, exist_ok=True)
    dst = PHOTO_DST / f"{product_id}.jpg"
    shutil.copyfile(src, dst)
    return dst.name


def main() -> int:
    for path in (CATALOG_XLSX, REGOS_XLSX, PHOTO_SRC):
        if not path.exists():
            print(f"Не найден путь: {path}")
            return 1

    try:
        prices = load_prices()
        items = load_catalog()
    except ImportError_ as e:
        print(f"Импорт остановлен. {e}")
        return 1

    if not items:
        print("В листе «Каталог MVP» нет позиций")
        return 1

    init_db()
    db = SessionLocal()
    added = {"категорий": 0, "товаров": 0, "фасовок": 0}
    photos_copied = 0

    categories: dict[str, Category] = {}
    for group in CATEGORY_ORDER:
        label = CATEGORY_LABEL[group]
        cat = db.scalar(select(Category).where(Category.name == label))
        if cat is None:
            cat = Category(name=label)
            db.add(cat)
            added["категорий"] += 1
        cat.sort_order = CATEGORY_ORDER.index(group)
        categories[group] = cat
    db.flush()

    touched: dict[int, list[str]] = {}   # товар -> коды его фасовок в этом файле

    for item in items:
        cat = categories.get(item["group"])
        if cat is None:
            print(f"  пропуск {item['code']}: неизвестная группа «{item['group']}»")
            continue

        # опознаём товар по коду REGOS — он стабилен, в отличие от названия
        variant = db.scalar(select(Variant).where(Variant.external_code == item["code"]))
        if variant is not None:
            product = variant.product
            product.name = item["name"]
            product.category_id = cat.id
        else:
            product = db.scalar(
                select(Product).where(Product.category_id == cat.id, Product.name == item["name"])
            )
            if product is None:
                product = Product(category_id=cat.id, name=item["name"])
                db.add(product)
                added["товаров"] += 1
            db.flush()
            variant = Variant(product_id=product.id, external_code=item["code"])
            db.add(variant)
            added["фасовок"] += 1

        price = prices.get(item["code"])
        variant.product_id = product.id
        variant.weight = item["weight"]
        variant.weight_grams = weight_to_grams(item["weight"])
        variant.price = price
        # фасовку без цены не показываем — иначе исказит «цену от» (BR-37)
        variant.is_active = price is not None

        db.flush()
        touched.setdefault(product.id, []).append(item["code"])

    # фото и публикация: пересчитываем целиком, чтобы удалённый снимок снимал товар с витрины
    for product_id, codes in touched.items():
        product = db.get(Product, product_id)
        photo = None
        for code in codes:
            if (PHOTO_SRC / f"p-{code}.jpg").exists():
                photo = copy_photo(code, product_id)
                photos_copied += 1
                break
        product.photo = photo
        # BR-34: публикуем карточку только когда есть и название, и фотография
        product.is_active = bool(product.name and product.photo)

    # позиции, исчезнувшие из файла, снимаем с продажи
    file_codes = {i["code"] for i in items}
    stale = [v for v in db.scalars(select(Variant).where(Variant.is_active))
             if v.external_code not in file_codes]
    for v in stale:
        v.is_active = False

    db.commit()

    # убираем снимки, на которые больше никто не ссылается (переименование, смена фото)
    in_use = {p.photo for p in db.scalars(select(Product)) if p.photo}
    removed = 0
    if PHOTO_DST.exists():
        for f in PHOTO_DST.iterdir():
            if f.is_file() and f.name not in in_use:
                f.unlink()
                removed += 1

    total_products = db.scalar(select(func.count()).select_from(Product))
    total_variants = db.scalar(select(func.count()).select_from(Variant))
    published = db.scalar(select(func.count()).select_from(Product).where(Product.is_active))

    print("Импорт завершён.")
    print(f"  создано: категорий {added['категорий']}, товаров {added['товаров']}, "
          f"фасовок {added['фасовок']}")
    print(f"  скопировано фотографий: {photos_copied}"
          + (f", удалено неиспользуемых: {removed}" if removed else ""))
    print(f"  всего в базе: товаров {total_products}, фасовок {total_variants}, "
          f"опубликовано {published}")

    no_price = db.scalar(select(func.count()).select_from(Variant).where(Variant.price.is_(None)))
    if no_price:
        print(f"\nБез цены (скрыты до получения прайс-листа): {no_price} фасовок")
    if stale:
        print(f"Исчезли из файла и сняты с продажи: {len(stale)} — "
              f"{', '.join(v.external_code for v in stale)}")

    # опубликованная карточка, у которой ни одной фасовки с ценой, покажется без цены
    priceless = [p for p in db.scalars(select(Product).where(Product.is_active))
                 if not any(v.is_active and v.price for v in p.variants)]
    if priceless:
        print(f"\nВНИМАНИЕ: опубликовано без единой цены — {len(priceless)}:")
        for p in priceless:
            print(f"     {p.name}")

    orphans = [p for p in db.scalars(select(Product)) if not p.variants]
    if orphans:
        print(f"\nВНИМАНИЕ: товары без фасовок — {len(orphans)}:")
        for p in orphans:
            print(f"     id={p.id} {p.name}")

    db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
